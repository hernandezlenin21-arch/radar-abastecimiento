"""Banco Mundial — Pink Sheet (CMO Historical Data Monthly). Datos abiertos, sin clave.

La URL del Excel cambia cada mes (lleva un hash), por eso se descubre leyendo
el enlace publicado en la página oficial de Commodity Markets. Puede fijarse
manualmente con la variable PINKSHEET_URL.
"""
from __future__ import annotations

import io
import os
import re
from functools import lru_cache

import openpyxl

from ..common import FetchResult, Observation, SourceError, http_get

LANDING = "https://www.worldbank.org/en/research/commodity-markets"
_LINK_RE = re.compile(r'https?://[^"\']+CMO-Historical-Data-Monthly\.xlsx', re.I)
_PERIOD_RE = re.compile(r"^(\d{4})M(\d{2})$")


def _clean(name) -> str:
    return re.sub(r"[\*\s]+", " ", str(name or "")).strip().lower()


def discover_url() -> str:
    override = os.environ.get("PINKSHEET_URL", "").strip()
    if override:
        return override
    html = http_get(LANDING).text
    m = _LINK_RE.search(html)
    if not m:
        raise SourceError("No se encontró el enlace CMO-Historical-Data-Monthly.xlsx en la página del Banco Mundial")
    return m.group(0)


def parse_workbook(content: bytes) -> dict[str, dict]:
    """Devuelve {nombre_limpio: {name, unit, obs:[Observation]}} de la hoja Monthly Prices."""
    wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    ws = next((wb[n] for n in wb.sheetnames if "monthly prices" in n.lower()), None)
    if ws is None:
        raise SourceError("El Excel del Banco Mundial no tiene la hoja 'Monthly Prices'")
    rows = list(ws.iter_rows(values_only=True))
    hdr_idx = next((i for i, r in enumerate(rows[:20]) if any(_clean(c).startswith("crude oil") for c in r)), None)
    if hdr_idx is None:
        raise SourceError("No se encontró la fila de encabezados en Monthly Prices")
    header, units = rows[hdr_idx], rows[hdr_idx + 1]
    out: dict[str, dict] = {}
    for col, name in enumerate(header):
        if col == 0 or not name:
            continue
        out[_clean(name)] = {"name": str(name).strip(), "unit": str(units[col] or "").strip(), "obs": []}
    cols = {col: _clean(name) for col, name in enumerate(header) if col and name}
    for r in rows[hdr_idx + 2:]:
        m = _PERIOD_RE.match(str(r[0] or "").strip())
        if not m:
            continue
        date = f"{m.group(1)}-{m.group(2)}-01"
        for col, key in cols.items():
            v = r[col] if col < len(r) else None
            if isinstance(v, (int, float)):
                out[key]["obs"].append(Observation(date, float(v)))
    return out


@lru_cache(maxsize=1)
def _download() -> tuple[str, dict]:
    url = discover_url()
    return url, parse_workbook(http_get(url, timeout=90).content)


def fetch(source_key: str, start: str) -> FetchResult:
    url, data = _download()
    key = _clean(source_key)
    match = data.get(key) or next((v for k, v in data.items() if k.startswith(key)), None)
    if match is None:
        raise SourceError(f"La columna '{source_key}' no existe en el Pink Sheet")
    obs = [o for o in match["obs"] if o.date >= start]
    unit = match["unit"].replace("$", "USD").replace("mt", "t") if match["unit"] else None
    return FetchResult(obs, source_url=url, unit=unit,
                       notes=[f"Columna original: {match['name']}"])
